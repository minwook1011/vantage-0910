"""데이터 허브 '매크로 × 섹터' — 일본 매크로 변수와 섹터별 대표 기업·업계 통계·지수의 상관관계 → docs/data/jp/macro_corr.json

2026-10-15 사용자: 매크로 변수마다 의류·외식·100엔숍 등 섹터의 대표 기업, 산업, 관련 지수가 어떻게 움직이는지 상관관계로 보고 싶다.
갱신은 예약 작업이 이틀에 한 번 pc_job.py 로 돌린다(이 스크립트는 받기·계산·저장만).

매크로(월간, 자동 수집):
  엔·달러 전년비(야후 JPY=X 월평균) · 도쿄 기온 전년차(기상청) · 방일 외국인 전년비 · 중국인 방일객 전년비(JNTO 엑셀) · 닛케이225 전년비(야후)
대상(월간):
  기업 매출 = docs/data/jp/monthly/<코드>.json 의 기존점 매출 전년比(없으면 전점) — 데이터 허브 월차 원본을 그대로 쓴다
  기업 주가 = 야후 <코드>.T 월평균 종가의 전년비
  섹터 평균 = 대표 기업 매출 전년비 평균 / 대표 기업 주가 전년비 평균(동일 가중 바스켓)
  업계 통계 = 일본푸드서비스협회 외식산업 시장동향(전체·패스트푸드·패밀리레스토랑·술집·카페 매출 전년비) — 협회 페이지에 최근 달만 걸려 있어
             지난 결과 파일의 값을 이어 붙여 쌓는다
  지수 = 닛케이225, 토픽스 소매업(TOPIX-17 소매 상장지수펀드 1630.T) 전년비
상관계수: 피어슨. 시차 0~6개월(매크로가 k개월 먼저). 겹치는 달이 10개 미만이면 비운다. 빈 값을 0으로 채우지 않는다.
한 소스가 실패하면 직전 결과 파일의 같은 계열을 그대로 쓴다.

사용: python build_jp_macro.py
"""
import calendar
import datetime as dt
import io
import json
import math
import re
import ssl
import sys
import time
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "data" / "jp" / "macro_corr.json"
MONTHLY = ROOT / "docs" / "data" / "jp" / "monthly"
KST = dt.timezone(dt.timedelta(hours=9))
UA = {"User-Agent": "Mozilla/5.0 (vantage jp macro)"}
START = "2019-01"
MAX_LAG = 6
MIN_N = 10

# 섹터와 대표 기업 — 월차 원본에 있는 회사만(시마무라·무인양품은 월차 자동 수집이 막혀 빠짐)
SECTORS = [
    {"id": "apparel", "label": "의류", "codes": ["9983", "7564", "2670", "3608", "8016", "3612", "8214", "8219", "7606"],
     "hint": "기온과 실질임금에 민감. 엔화 약세는 원가 부담(수입 의존)."},
    {"id": "food", "label": "외식", "codes": ["7550", "2702", "3197", "7581", "3397", "9861", "9887", "2695", "3087", "7611"],
     "hint": "인건비·식재료(엔화) 원가가 핵심. 수요는 비교적 단단하고 관광지 매장은 방일객 영향."},
    {"id": "yen100", "label": "100엔숍", "codes": ["2782", "2698"],
     "hint": "판매가 고정이라 엔화 약세가 이익률을 직접 깎는다. 수요는 절약 소비로 오히려 늘기 쉽다."},
    {"id": "dept", "label": "백화점", "codes": ["3099", "8233", "3086", "8237", "8244"],
     "hint": "방일객(면세)과 주가(부유층 자산 효과)에 가장 민감."},
    {"id": "drug", "label": "드럭스토어", "codes": ["3088", "7649", "9267", "3549", "7679"],
     "hint": "생활필수품 중심이라 경기에 둔감. 도심 매장은 방일객 화장품 수요."},
]
INDUSTRY_ROWS = [  # 외식협회 월간 파일의 행 이름(앞부분) → 우리 id
    ("jf_all", "외식 전체", ("全", "体")),
    ("jf_ff", "패스트푸드", None),
    ("jf_fr", "패밀리레스토랑", None),
    ("jf_pub", "술집·주점", None),
    ("jf_cafe", "카페", ("喫茶",)),
]


def get(url, tries=3):
    last = None
    for i in range(tries):
        for ctx in (None, ssl._create_unverified_context()):  # 이 PC는 SSL 가로채기가 있어 공개 통계만 검증 없이 재시도
            try:
                with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=40, context=ctx) as r:
                    return r.read()
            except urllib.error.HTTPError as e:
                if e.code == 404:
                    raise
                last = e
            except Exception as e:  # noqa: BLE001
                last = e
        time.sleep(2 + i * 3)
    raise last


def ym(y, m):
    return f"{y:04d}-{m:02d}"


def shift(k, n):
    y, m = int(k[:4]), int(k[5:7]) + n
    while m <= 0:
        y, m = y - 1, m + 12
    while m > 12:
        y, m = y + 1, m - 12
    return ym(y, m)


def this_month():
    n = dt.datetime.now(KST)
    return ym(n.year, n.month)


def yoy_pct(level):
    return {k: round((v / level[shift(k, -12)] - 1) * 100, 2) for k, v in level.items() if level.get(shift(k, -12))}


def yahoo_monthly(sym):
    j = json.loads(get(f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?range=10y&interval=1d"))
    r = j["chart"]["result"][0]
    acc = {}
    for t, c in zip(r["timestamp"], r["indicators"]["quote"][0]["close"]):
        if c is None:
            continue
        d = dt.datetime.fromtimestamp(t, KST)
        acc.setdefault(ym(d.year, d.month), []).append(c)
    cur = this_month()
    return {k: sum(v) / len(v) for k, v in acc.items() if k != cur and k >= shift(START, -12)}


def fetch_temp():
    temp = {}
    now = dt.datetime.now(KST)
    for y in range(int(START[:4]) - 1, now.year + 1):
        html = get(f"https://www.data.jma.go.jp/stats/etrn/view/monthly_s1.php?prec_no=44&block_no=47662&year={y}&month=&day=&view=").decode("utf-8", "replace")
        for row in re.findall(r'<tr class="mtx"[^>]*>(.*?)</tr>', html, re.S):
            cells = [re.sub(r"<[^>]+>", "", c).strip() for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, re.S)]
            if len(cells) > 7 and cells[0].isdigit():
                k = ym(y, int(cells[0]))
                v = re.sub(r"[^\d.\-]", "", cells[7])
                if k < this_month() and v not in ("", "-", "."):
                    temp[k] = float(v)
        time.sleep(1)
    return {k: round(v - temp[shift(k, -12)], 1) for k, v in temp.items() if shift(k, -12) in temp}


def fetch_inbound():
    import openpyxl
    page = get("https://www.jnto.go.jp/statistics/data/visitors-statistics/").decode("utf-8", "replace")
    link = re.search(r'href="([^"]+\.xlsx)"', page).group(1)
    if link.startswith("/"):
        link = "https://www.jnto.go.jp" + link
    wb = openpyxl.load_workbook(io.BytesIO(get(link)), read_only=True, data_only=True)
    total, china = {}, {}
    for name in wb.sheetnames:
        if not re.fullmatch(r"\d{4}", name) or int(name) < int(START[:4]) - 1:
            continue
        y = int(name)
        head = None
        for r in wb[name].iter_rows(values_only=True):
            cells = list(r)
            if head is None and any(isinstance(c, str) and c.strip() == "1月" for c in cells):
                head = {i: int(c.strip()[:-1]) for i, c in enumerate(cells) if isinstance(c, str) and re.fullmatch(r"\d{1,2}月", c.strip())}
                continue
            if head and isinstance(cells[0], str) and cells[0].strip() in ("総数", "中国"):
                dst = total if cells[0].strip() == "総数" else china
                for i, m in head.items():
                    if isinstance(cells[i], (int, float)):
                        dst[ym(y, m)] = float(cells[i])
    return yoy_pct(total), yoy_pct(china), link


def parse_jf(raw):
    """외식협회 월간 파일 — '全店' 시트에서 업태 합계 행의 매출 전년比(1.048 → +4.8%)"""
    import xlrd
    s = xlrd.open_workbook(file_contents=raw).sheet_by_index(0)
    out, totals = {}, []
    for i in range(s.nrows):
        row = [str(x).strip() for x in s.row_values(i)]
        label = "".join(row[:6]).replace("　", "").replace(" ", "")
        try:
            v = float(row[8])
        except (ValueError, IndexError):
            continue
        if label == "全体":
            out["jf_all"] = v
        elif label == "合計":
            totals.append(v)
        elif label.startswith("喫茶"):
            out["jf_cafe"] = v
    for k, v in zip(("jf_ff", "jf_fr", "jf_pub"), totals):  # 합계 행은 패스트푸드 → 패밀리레스토랑 → 술집 순서
        out[k] = v
    return {k: round((v - 1) * 100, 2) for k, v in out.items()}


def fetch_food_service():
    page = get("https://www.jfnet.or.jp/data/data_c.html").decode("utf-8", "replace")
    links = set(re.findall(r'href="(https://www\.jfnet\.or\.jp/wp/wp-content/uploads/[^"]+data_c_m(\d{4})_(\d{1,2})[^"/]*\.xlsx?)"', page))
    # 페이지에서 내려간 달은 올린 달(다음 달·다다음 달) 규칙으로 주소를 추정한다(2025-07 이후 파일만 이 규칙이 맞았다)
    have = {ym(int(y), int(m)) for _, y, m in links}
    k = "2025-07"
    while k < this_month():
        if k not in have:
            for up in (1, 2):
                u = shift(k, up)
                url = f"https://www.jfnet.or.jp/wp/wp-content/uploads/{u[:4]}/{u[5:]}/data_c_m{k[:4]}_{k[5:]}.xls"
                try:
                    urllib.request.urlopen(urllib.request.Request(url, method="HEAD", headers=UA), timeout=15, context=ssl._create_unverified_context())
                    links.add((url, k[:4], k[5:]))
                    break
                except Exception:  # noqa: BLE001
                    pass
        k = shift(k, 1)
    res = {}
    for url, y, m in links:
        try:
            for k, v in parse_jf(get(url)).items():
                res.setdefault(k, {})[ym(int(y), int(m))] = v
        except Exception as e:  # noqa: BLE001
            print(f"[warn] 외식협회 {y}-{m}: {e}", file=sys.stderr)
        time.sleep(0.5)
    return res


def company_sales(code):
    p = MONTHLY / f"{code}.json"
    if not p.exists():
        return None, None
    j = json.loads(p.read_text(encoding="utf-8"))
    cols = j.get("cols") or []
    pick = next((c for c in cols if c.startswith("既存店") and "売上" in c and "전년比" in c), None) \
        or next((c for c in cols if "売上" in c and "전년比" in c), None)
    if not pick:
        return None, None
    s = {k: round(v - 100, 2) for k, v in (j.get("data") or {}).items() if isinstance(v.get(pick), (int, float)) for v in [v[pick]]}
    return s, pick


def pearson(xs, ys):
    n = len(xs)
    if n < MIN_N:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    sxx = sum((x - mx) ** 2 for x in xs)
    syy = sum((y - my) ** 2 for y in ys)
    if sxx == 0 or syy == 0:
        return None
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    return sxy / math.sqrt(sxx * syy), sxy / sxx, n


def corr(macro, target):
    """lags[k] = 매크로를 k개월 앞당겨 맞춘 상관계수. slope = 매크로 1단위당 대상 변화(시차 0)"""
    lags, info = [], {}
    for k in range(MAX_LAG + 1):
        pairs = [(macro[shift(t, -k)], v) for t, v in target.items() if shift(t, -k) in macro]
        r = pearson([p[0] for p in pairs], [p[1] for p in pairs])
        lags.append(None if r is None else round(r[0], 3))
        if k == 0 and r:
            info = {"slope": round(r[1], 3), "n": r[2]}
    if all(v is None for v in lags):
        return None
    best = max((k for k, v in enumerate(lags) if v is not None), key=lambda k: abs(lags[k]))
    return {"r0": lags[0], "lags": lags, "best": best, "rb": lags[best], **info}


def clip(s):
    return {k: v for k, v in sorted(s.items()) if k >= START}


def main():
    prev = json.loads(OUT.read_text(encoding="utf-8")) if OUT.exists() else {}
    prev_series = prev.get("series", {})
    errors = []
    names = {it["c"]: it.get("ko") or it.get("n") for it in json.loads((MONTHLY / "index.json").read_text(encoding="utf-8")).get("items", [])}

    series, meta = {}, {}

    def keep(sid, fn, **m):
        try:
            v = fn()
            if not v:
                raise ValueError("빈 결과")
            series[sid] = clip(v)
        except Exception as e:  # noqa: BLE001
            errors.append(f"{sid}: {e}")
            print(f"[fail] {sid}: {e}", file=sys.stderr)
            if sid in prev_series:
                series[sid] = prev_series[sid]
        meta[sid] = m

    # 1) 매크로
    macros = [
        {"id": "m_usdjpy", "label": "엔·달러 전년비", "unit": "%", "src": "야후 파이낸스(JPY=X 월평균)", "hint": "+면 엔화 약세. 수입 원가가 오르고 방일객은 늘어난다."},
        {"id": "m_temp", "label": "도쿄 기온 전년차", "unit": "℃", "src": "기상청(도쿄 월평균 기온)", "hint": "+면 작년 같은 달보다 따뜻함. 가을·겨울엔 두꺼운 옷 판매에 불리."},
        {"id": "m_wxfit", "label": "옷 팔기 좋은 날씨", "unit": "℃", "src": "기상청(도쿄 월평균 기온 전년차 × 계절 부호)", "hint": "3~8월은 작년보다 더우면 +, 9~2월은 작년보다 추우면 +. 이번 여름처럼 서늘한 여름은 −."},
        {"id": "m_inbound", "label": "방일 외국인 전년비", "unit": "%", "src": "일본정부관광국(JNTO) 월별 방일객 수", "hint": "백화점 면세·도심 드럭스토어·관광지 외식."},
        {"id": "m_china", "label": "중국인 방일객 전년비", "unit": "%", "src": "일본정부관광국(JNTO) 월별 방일객 수", "hint": "2025년 말 중국 정부의 일본 여행 자제 권고 이후 급감."},
        {"id": "m_nikkei", "label": "닛케이225 전년비", "unit": "%", "src": "야후 파이낸스(^N225 월평균)", "hint": "주가가 오르면 부유층·고령층 소비(자산 효과)."},
    ]
    keep("m_usdjpy", lambda: yoy_pct(yahoo_monthly("JPY=X")))
    keep("m_temp", fetch_temp)
    if "m_temp" in series:  # 여름옷은 더워야, 가을·겨울옷은 추워야 팔린다 — 계절 부호를 곱한다
        series["m_wxfit"] = {k: round(v if 3 <= int(k[5:7]) <= 8 else -v, 1) for k, v in series["m_temp"].items()}
    jnto_link = ""
    try:
        inb, chn, jnto_link = fetch_inbound()
        series["m_inbound"], series["m_china"] = clip(inb), clip(chn)
    except Exception as e:  # noqa: BLE001
        errors.append(f"m_inbound: {e}")
        for sid in ("m_inbound", "m_china"):
            if sid in prev_series:
                series[sid] = prev_series[sid]
    nikkei = {}
    try:
        nikkei = yoy_pct(yahoo_monthly("^N225"))
        series["m_nikkei"] = clip(nikkei)
    except Exception as e:  # noqa: BLE001
        errors.append(f"m_nikkei: {e}")
        if "m_nikkei" in prev_series:
            series["m_nikkei"] = prev_series["m_nikkei"]

    # 2) 지수
    indexes = [{"id": "ix_nikkei", "label": "닛케이225", "src": "야후 파이낸스(^N225)"},
               {"id": "ix_retail", "label": "토픽스 소매업", "src": "야후 파이낸스(1630.T, TOPIX-17 소매 상장지수펀드)"}]
    if "m_nikkei" in series:
        series["ix_nikkei"] = series["m_nikkei"]
    retail_lv = {}

    def _retail():
        retail_lv.update(yahoo_monthly("1630.T"))
        return yoy_pct(retail_lv)
    keep("ix_retail", _retail)
    if retail_lv:
        series["il_retail"] = {k: round(v, 1) for k, v in clip(retail_lv).items()}

    # 3) 업계 통계(외식) — 협회 페이지엔 최근 달만 있어 지난 값을 이어 붙인다
    jf_new = {}
    try:
        jf_new = fetch_food_service()
        if not jf_new:
            raise ValueError("파일 없음")
    except Exception as e:  # noqa: BLE001
        errors.append(f"외식협회: {e}")
    for sid, label, _ in INDUSTRY_ROWS:
        merged = dict(prev_series.get(sid, {}))
        merged.update(jf_new.get(sid, {}))
        if merged:
            series[sid] = clip(merged)

    # 4) 기업 매출·주가, 섹터 평균
    sectors_out = []
    for sec in SECTORS:
        rows = []
        sales_all, px_all = [], []
        for code in sec["codes"]:
            sales, col = company_sales(code)
            row = {"code": code, "name": names.get(code, code), "sales": None, "px": None, "sales_col": col}
            if sales:
                sid = f"s_{code}"
                series[sid] = clip(sales)
                row["sales"] = sid
                sales_all.append(series[sid])
            pid = f"p_{code}"
            try:
                lv = yahoo_monthly(f"{code}.T")
                series[pid] = clip(yoy_pct(lv))
                series[f"pl_{code}"] = {k: round(v, 1) for k, v in clip(lv).items()}
                row["px"] = pid
            except Exception as e:  # noqa: BLE001
                errors.append(f"{pid}: {e}")
                if pid in prev_series:
                    series[pid] = prev_series[pid]
                    row["px"] = pid
            if row["px"]:
                px_all.append(series[pid])
            rows.append(row)
            time.sleep(0.4)

        def avg(lst):
            keys = sorted(set().union(*lst)) if lst else []
            out = {}
            for k in keys:
                vals = [s[k] for s in lst if k in s]
                if len(vals) >= max(1, len(lst) // 2):  # 절반 이상 값이 있는 달만
                    out[k] = round(sum(vals) / len(vals), 2)
            return out
        agg = []
        if sales_all:
            series[f"sa_{sec['id']}"] = avg(sales_all)
            agg.append({"id": f"sa_{sec['id']}", "kind": "sales", "label": f"{sec['label']} 대표 기업 매출 평균"})
        if px_all:
            series[f"pa_{sec['id']}"] = avg(px_all)
            agg.append({"id": f"pa_{sec['id']}", "kind": "px", "label": f"{sec['label']} 대표 기업 주가 바스켓"})
        industry = [{"id": sid, "label": label, "src": "일본푸드서비스협회 외식산업 시장동향(매출 전년비, 전 점포)"} for sid, label, _ in INDUSTRY_ROWS if sec["id"] == "food" and sid in series]
        sectors_out.append({"id": sec["id"], "label": sec["label"], "hint": sec["hint"], "companies": rows, "aggregates": agg, "industry": industry})

    # 5) 상관계수: 매크로 × 모든 대상
    targets = [sid for sid in series if sid.split("_")[0] in ("s", "p", "sa", "pa", "jf", "ix")]
    corr_out = {}
    for m in macros:
        if m["id"] not in series:
            continue
        corr_out[m["id"]] = {}
        for t in targets:
            c = corr(series[m["id"]], series[t])
            if c:
                corr_out[m["id"]][t] = c

    out = {"schema_version": 1, "generated_at": dt.datetime.now(KST).isoformat(timespec="seconds"),
           "start": START, "max_lag": MAX_LAG, "min_n": MIN_N,
           "macros": [m for m in macros if m["id"] in series], "indexes": [i for i in indexes if i["id"] in series],
           "sectors": sectors_out, "series": series, "corr": corr_out, "errors": errors,
           "sources": {"jnto": jnto_link, "jfnet": "https://www.jfnet.or.jp/data/data_c.html", "jma": "https://www.data.jma.go.jp/stats/etrn/"}}
    OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    n_corr = sum(len(v) for v in corr_out.values())
    print(f"저장 · 계열 {len(series)}개 · 상관계수 {n_corr}개 · 실패 {len(errors)}건")


if __name__ == "__main__":
    main()
