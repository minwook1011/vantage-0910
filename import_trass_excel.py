"""사용자가 관리하는 TRASS 기준 엑셀(메모리_잠정치_합산+비중_*.xlsx)의 '전세계' 월별 금액·중량을
docs/data/trass_exports.json 에 월 전체(M) 관측값으로 넣는다.

엑셀 구조 (품목 시트마다 같음 · 2026-09 확인)
  C3        HS 코드
  행 40~51  1~12월,  D~J열 = 20~26년 금액(달러),  V~AB열 = 20~26년 순중량(kg)

진행 중인 달(오늘이 속한 달)은 1~10일·1~20일 누적 잠정치라 월 전체(M)로 넣지 않는다.
그 달은 TRASS 순별 입력(ingest_trass.py)의 D10/D20 관측값이 담당한다.
엑셀은 공개 저장소에 올리지 않는다 — 이 스크립트가 숫자만 옮긴다.

    python import_trass_excel.py "C:/.../수출입/메모리_잠정치_합산+비중_20260730.xlsx"
"""
import os
import sys
from datetime import datetime, timedelta, timezone

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ingest_trass as T  # noqa: E402  같은 upsert(정정 이력 포함)를 쓴다

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

KST = timezone(timedelta(hours=9))
SHEETS = {  # 시트 → (품목 id, 표시 이름, 설명)
    "잠정_메모리": ("memory", "메모리", "HS 8542.32 메모리 집적회로 전체"),
    "잠정_DRAM": ("dram", "DRAM", "모노리식 DRAM. HBM·DRAM 모듈 전체 출하가 아님"),
    "잠정_Flash": ("flash", "Flash 메모리", "플래시 메모리. 순수 NAND 전용으로 단정하지 않음"),
    "잠정_MCP": ("mcp", "MCP (HBM 포함 추정)", "복합구조칩(MCP). 시장은 HBM이 주로 여기 잡힌다고 보지만 HS 명칭상 HBM 전용은 아님"),
    "잠정_DRAM모듈": ("dram_module", "DRAM 모듈", "HS 8473.30-4060 컴퓨터용 DRAM 모듈"),
    "잠정_SSD": ("ssd", "SSD", "HS 8523.51-1000 SSD (반도체 저장장치)"),
    "잠정_MCOs": ("mcos", "MCOs", "HS 8542.32-4000 다중부품 집적회로(MCO)"),
    "잠정_프로세서": ("processor", "프로세서·컨트롤러", "HS 8542.31-1000 프로세서·컨트롤러"),
    "잠정_기타IC": ("other_ic", "기타 IC", "HS 8542.39-1000 기타 집적회로"),
}
YEARS = [2020, 2021, 2022, 2023, 2024, 2025, 2026]  # D~J / V~AB 열 순서


def ensure_item(data, item_id, label, hs, note):
    for it in data["items"]:
        if it["id"] == item_id:
            if hs and not it.get("hs"):
                it["hs"], it["verify"] = hs, False
                it["hs_note"] = note
            return
    # 기존 반도체 품목 뒤에 끼워 넣는다
    idx = max((i for i, it in enumerate(data["items"]) if it.get("group") == "반도체"), default=len(data["items"]) - 1) + 1
    data["items"].insert(idx, {"id": item_id, "label": label, "group": "반도체", "hs": hs, "hs_note": note})


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    data = T.load()
    this_month = datetime.now(KST).strftime("%Y-%m")
    source = "TRASS 월간 (사용자 기준 엑셀 " + os.path.basename(path) + ")"
    counts = {"added": 0, "revised": 0, "unchanged": 0}
    for sheet, (item_id, label, note) in SHEETS.items():
        if sheet not in wb.sheetnames:
            print(f"  [skip] 시트 없음: {sheet}")
            continue
        ws = wb[sheet]
        hs = ws["C3"].value
        hs = str(int(hs)) if isinstance(hs, (int, float)) else (str(hs).strip() if hs else None)
        ensure_item(data, item_id, label, hs, note)
        rows = list(ws.iter_rows(min_row=40, max_row=51, max_col=28, values_only=True))
        for m_idx, r in enumerate(rows):
            for y_idx, year in enumerate(YEARS):
                usd, kg = r[3 + y_idx], r[21 + y_idx]
                month = f"{year}-{m_idx + 1:02d}"
                if not isinstance(usd, (int, float)) or usd <= 0 or month >= this_month:
                    continue  # 빈 칸 · 진행 중인 달(순별 누적 잠정치)은 넣지 않는다
                kg = float(kg) if isinstance(kg, (int, float)) and kg > 0 else None
                counts[T.upsert(data, item_id, month, "M", float(usd), kg, source, None)] += 1
    T.save(data)
    print(f"완료 — 추가 {counts['added']}, 정정 {counts['revised']}, 변화없음 {counts['unchanged']} · 진행 중인 {this_month} 은 제외")


if __name__ == "__main__":
    main()
